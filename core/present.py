"""発表用グラフ（静止画）と、津山の発表デザインのPPTXを作る。

図のスタイルは発表スライドで使った図と同じ: 白背景・左寄せの太字見出し（結論を書く）・控えめな罫線・
値ラベル・脚注に人数と検定結果。色は役割で固定（カテゴリは決まった順、循環させない）。
"""
import io
import shutil
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

from .analyze import K_MIN  # noqa: E402

BLUE, ORANGE, AQUA, YELLOW, MAGENTA, GREEN, VIOLET, RED = (
    "#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948")
CAT = [BLUE, ORANGE, AQUA, YELLOW, MAGENTA, GREEN, VIOLET, RED]
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e4e3df"
AGE_ORDER = ["〜29", "30代", "40代", "50代", "60-64", "65-74", "75+"]
TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "tsuyama_template.pptx"

def _jp_font():
    """このPCにある日本語フォントを1つ選ぶ（Mac: ヒラギノ／Windows: 游ゴシック・メイリオ）。"""
    from matplotlib import font_manager as fm
    have = {f.name for f in fm.fontManager.ttflist}
    for name in ["Hiragino Sans", "YuGothic", "Yu Gothic", "BIZ UDPGothic", "Meiryo", "Noto Sans JP", "Noto Sans CJK JP"]:
        if name in have:
            return name
    return "sans-serif"


plt.rcParams.update({
    "font.family": _jp_font(),
    "font.size": 10, "axes.edgecolor": GRID, "axes.labelcolor": INK2, "axes.titlesize": 12,
    "axes.titleweight": "bold", "axes.titlecolor": INK, "axes.titlelocation": "left",
    "xtick.color": INK2, "ytick.color": INK2, "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6,
    "axes.axisbelow": True, "axes.spines.top": False, "axes.spines.right": False, "legend.frameon": False,
})


def fmt(v):
    if pd.isna(v):
        return "-"
    if float(v).is_integer():
        return f"{int(v):,}"
    a = abs(v)
    return f"{v:,.0f}" if a >= 100 else f"{v:.1f}" if a >= 10 else f"{v:.2f}"


def _cat_num(r):
    return (r["col1"], r["col2"], r["変数1"], r["変数2"]) if r["kind1"] == "cat" else (r["col2"], r["col1"], r["変数2"], r["変数1"])


def headline(df, r):
    """図の見出し（結論）を自動で作る。"""
    if r["kind1"] == "num" and r["kind2"] == "num":
        return f"「{r['変数1']}」が高い人ほど、「{r['変数2']}」が{'高い' if r['効果量'] > 0 else '低い'}"
    if r["kind1"] == "cat" and r["kind2"] == "cat":
        return f"「{r['変数1']}」によって、「{r['変数2']}」の内訳が異なる"
    cat, num, lc, ln = _cat_num(r)
    meds = df[[cat, num]].dropna().groupby(cat)[num].agg(["median", "size"])
    meds = meds[meds["size"] >= K_MIN]["median"].sort_values(ascending=False)
    if len(meds) == 2:
        return f"「{lc}：{meds.index[0]}」の人は「{ln}」が高い（{fmt(meds.iloc[0])} vs {meds.index[1]} {fmt(meds.iloc[1])}）"
    return f"「{ln}」は{lc}によって異なる（最も高い：{meds.index[0]}、最も低い：{meds.index[-1]}）"


def footnote(r):
    q = "q<0.001" if r["q"] < 0.001 else f"q={r['q']:.3f}"
    kind = {"ρ": "Spearman相関", "r": "Mann-Whitney検定", "ε²": "Kruskal-Wallis検定", "V": "カイ二乗検定"}[r["効果量の種類"]]
    base = f"{r['データ']}　n={r['n']:,}。{kind}（{r['効果量の種類']}={r['効果量']:+.2f}、{q}、多重比較補正後）。"
    extra = {2: "横軸を5等分した区間ごとの中央値、縦線は四分位範囲。",
             1: "棒は中央値。10人未満のカテゴリは非表示。",
             0: "各カテゴリの中での内訳。10人未満のカテゴリは非表示。"}[(r["kind1"] == "num") + (r["kind2"] == "num")]
    return base + extra


def make_figure(df, r, title=None, style="区間ごとの中央値"):
    """matplotlib の Figure を返す。title=None なら見出しなし（スライド側にタイトルを置く場合）。"""
    if r["kind1"] == "num" and r["kind2"] == "num":
        c1, c2, l1, l2 = r["col1"], r["col2"], r["変数1"], r["変数2"]
        m = df[[c1, c2]].dropna()
        fig, ax = plt.subplots(figsize=(8, 4.4))
        if style == "散布図＋傾向線":
            s = m.sample(min(len(m), 3000), random_state=0)
            ax.scatter(s[c1], s[c2], s=10, color=BLUE, alpha=0.3, linewidths=0)
            m["bin"] = pd.qcut(m[c1].rank(method="first"), 10, labels=False)
            t = m.groupby("bin").agg(x=(c1, "median"), y=(c2, "median"))
            ax.plot(t.x, t.y, color=ORANGE, lw=2.5, marker="o", ms=5, label="10分位ごとの中央値")
            for a_, c in (("x", c1), ("y", c2)):
                lo, hi = m[c].quantile([0.01, 0.99])
                pad = (hi - lo) * 0.05 or 1
                (ax.set_xlim if a_ == "x" else ax.set_ylim)(lo - pad, hi + pad)
            ax.legend(loc="upper left")
            ax.set_xlabel(l1)
        else:
            m["bin"] = pd.qcut(m[c1].rank(method="first"), 5, labels=False)
            t = m.groupby("bin").agg(lo=(c1, "min"), hi=(c1, "max"), y=(c2, "median"),
                                     q1=(c2, lambda s: s.quantile(0.25)), q3=(c2, lambda s: s.quantile(0.75)), n=(c2, "size"))
            xs = np.arange(len(t))
            colors = [ORANGE if r["効果量"] > 0 and i == len(t) - 1 or r["効果量"] < 0 and i == 0 else BLUE for i in xs]
            ax.bar(xs, t.y, color=colors, width=0.62)
            ax.errorbar(xs, t.y, yerr=[t.y - t.q1, t.q3 - t.y], fmt="none", ecolor=INK2, capsize=3, lw=1)
            for x, y in zip(xs, t.y):
                ax.text(x + 0.33, y, fmt(y), va="center", ha="left", fontsize=9, color=INK, fontweight="bold")
            ax.set_xticks(xs, [f"{['低い', 'やや低い', '中間', 'やや高い', '高い'][i]}\n{fmt(a)}〜{fmt(b)}" for i, (a, b) in enumerate(zip(t.lo, t.hi))],
                          fontsize=8.5)
            ax.set_xlabel(f"{l1}（人数が等しい5区間）")
            ax.grid(axis="x", visible=False)
        ax.set_ylabel(l2)
    elif r["kind1"] == "cat" and r["kind2"] == "cat":
        c1, c2, l1, l2 = r["col1"], r["col2"], r["変数1"], r["変数2"]
        m = df[[c1, c2]].dropna()
        k1, k2 = m[c1].value_counts(), m[c2].value_counts()
        m = m[m[c1].isin(k1[k1 >= K_MIN].index) & m[c2].isin(k2[k2 >= K_MIN].index)]
        ct = pd.crosstab(m[c1], m[c2])
        rows = [o for o in AGE_ORDER if o in ct.index] or list(ct.index)
        if len(rows) != len(ct.index):
            rows = list(ct.index)
        ct = ct.loc[rows]
        pct = ct.div(ct.sum(axis=1), axis=0) * 100
        fig, ax = plt.subplots(figsize=(8, max(2.6, 0.55 * len(pct) + 1.4)))
        left = np.zeros(len(pct))
        y = np.arange(len(pct))
        for i, col in enumerate(pct.columns):
            c = CAT[i] if len(pct.columns) <= len(CAT) else BLUE
            ax.barh(y, pct[col], left=left, color=c, height=0.62, label=str(col), edgecolor="white", linewidth=2)
            for yy, (l, v) in enumerate(zip(left, pct[col])):
                if v >= 6:
                    ax.text(l + v / 2, yy, f"{v:.0f}%", ha="center", va="center", fontsize=9, color="white", fontweight="bold")
            left += pct[col].values
        ax.set_yticks(y, [f"{i}（n={ct.loc[i].sum():,}）" for i in pct.index])
        ax.invert_yaxis()
        ax.set_xlim(0, 100)
        ax.set_xlabel(f"「{l2}」の内訳（%）")
        ax.legend(ncol=min(len(pct.columns), 5), loc="lower center", bbox_to_anchor=(0.5, 1.0), fontsize=9)
        ax.grid(axis="y", visible=False)
    else:
        cat, num, lc, ln = _cat_num(r)
        m = df[[cat, num]].dropna()
        g = m.groupby(cat)[num].agg(["median", "size"])
        g = g[g["size"] >= K_MIN]
        natural = [o for o in AGE_ORDER if o in g.index]
        g = g.loc[natural] if len(natural) == len(g) else g.sort_values("median", ascending=False)
        hi, lo = g["median"].idxmax(), g["median"].idxmin()
        colors = [ORANGE if i == hi else BLUE if i == lo else "#9ec5f4" for i in g.index] if len(g) > 2 else [ORANGE, BLUE] if g["median"].iloc[0] >= g["median"].iloc[1] else [BLUE, ORANGE]
        labels = [f"{i}\n(n={int(s):,})" for i, s in zip(g.index, g["size"])]
        if len(g) <= 5:
            fig, ax = plt.subplots(figsize=(8, 4.2))
            xs = np.arange(len(g))
            ax.bar(xs, g["median"], color=colors, width=0.6)
            for x, v in zip(xs, g["median"]):
                ax.text(x, v, fmt(v), ha="center", va="bottom", fontsize=10, color=INK, fontweight="bold")
            ax.set_xticks(xs, labels, fontsize=9)
            ax.set_ylabel(f"{ln}（中央値）")
            ax.set_xlabel(lc)
            ax.grid(axis="x", visible=False)
            ax.set_ylim(0, g["median"].max() * 1.15 if g["median"].max() > 0 else None)
        else:
            fig, ax = plt.subplots(figsize=(8, 0.42 * len(g) + 1.4))
            ys = np.arange(len(g))
            ax.barh(ys, g["median"], color=colors, height=0.62)
            for yv, v in zip(ys, g["median"]):
                ax.text(v, yv, " " + fmt(v), va="center", fontsize=9, color=INK, fontweight="bold")
            ax.set_yticks(ys, [lab.replace("\n", " ") for lab in labels], fontsize=9)
            ax.invert_yaxis()
            ax.set_xlabel(f"{ln}（中央値）")
            ax.grid(axis="y", visible=False)
    if title:
        fig.suptitle(title, x=0.01, ha="left", fontweight="bold", color=INK, fontsize=12.5)
    fig.text(0.01, -0.02, footnote(r), fontsize=8, color=INK2, ha="left", va="top", wrap=True)
    fig.tight_layout()
    return fig


def to_png(fig, dpi=200) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


# ================================================================ PPTX（津山の発表デザイン）
def build_pptx(items, subtitle="組み合わせ分析") -> bytes:
    """items: [{"png": bytes, "title": str, "kicker": str, "note": str}]。テンプレートがあれば津山デザインで作る。"""
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    FONT = "Yu Gothic"

    def ja(run, size, color, bold=False):
        f = run.font
        f.name, f.size, f.bold, f.color.rgb = FONT, Pt(size), bold, color
        rpr = run._r.get_or_add_rPr()
        rpr.set("lang", "ja-JP")
        rpr.set("altLang", "en-US")
        for tag in ("a:latin", "a:ea", "a:cs"):
            e = rpr.find(qn(tag))
            if e is None:
                e = rpr.makeelement(qn(tag), {})
                rpr.append(e)
            e.set("typeface", FONT)

    use_tpl = TEMPLATE.exists()
    if use_tpl:
        tmp = Path(tempfile.mkdtemp()) / "t.pptx"
        shutil.copyfile(TEMPLATE, tmp)
        prs = Presentation(tmp)
        lst = prs.slides._sldIdLst
        for sid in list(lst):
            prs.part.drop_rel(sid.rId)
            lst.remove(sid)
        layout = prs.slide_layouts[11] if len(prs.slide_layouts) > 11 else prs.slide_layouts[6]
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        layout = prs.slide_layouts[6]
    W, H = prs.slide_width / 914400, prs.slide_height / 914400
    PINK, INKC = RGBColor(0xB6, 0x37, 0x5F), RGBColor(0, 0, 0)

    for i, it in enumerate(items, 1):
        s = prs.slides.add_slide(layout)
        placed_title = placed_kicker = False
        for ph in list(s.placeholders):
            idx = ph.placeholder_format.idx
            if idx == 0:
                ph.text_frame.text = ""
                ja(ph.text_frame.paragraphs[0].add_run(), 22, INKC, True)
                ph.text_frame.paragraphs[0].runs[0].text = it["title"]
                placed_title = True
            elif idx == 13:
                ph.text_frame.text = ""
                ja(ph.text_frame.paragraphs[0].add_run(), 14, PINK, True)
                ph.text_frame.paragraphs[0].runs[0].text = it["kicker"]
                placed_kicker = True
            elif idx == 14:
                ph._element.getparent().remove(ph._element)
        if not placed_title:
            tb = s.shapes.add_textbox(Inches(0.3), Inches(0.25), Inches(W - 0.6), Inches(1.0))
            tb.text_frame.word_wrap = True
            ja(tb.text_frame.paragraphs[0].add_run(), 22, INKC, True)
            tb.text_frame.paragraphs[0].runs[0].text = it["title"]
        if not placed_kicker:
            tb = s.shapes.add_textbox(Inches(0.3), Inches(1.3), Inches(W - 0.6), Inches(0.4))
            ja(tb.text_frame.paragraphs[0].add_run(), 14, PINK, True)
            tb.text_frame.paragraphs[0].runs[0].text = it["kicker"]
        # 図を白パネル内に縦横比を保って配置
        x0, y0, bw, bh = 0.42, 1.95, W - 0.84, H - 2.75
        iw, ih = Image.open(io.BytesIO(it["png"])).size
        sc = min(bw / iw, bh / ih)
        pw, ph_ = iw * sc, ih * sc
        s.shapes.add_picture(io.BytesIO(it["png"]), Inches(x0 + (bw - pw) / 2), Inches(y0 + (bh - ph_) / 2), Inches(pw), Inches(ph_))
        nb = s.shapes.add_textbox(Inches(W - 0.75), Inches(H - 0.43), Inches(0.6), Inches(0.35))
        p = nb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        ja(p.add_run(), 9, RGBColor(0x40, 0x40, 0x40))
        p.runs[0].text = str(i)
        nb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
